import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

A = dict(  # key: (label, value, note)
 pro_m=("Pro сарын үнэ (НӨАТ орсон), ₮", 9900, "plans хүснэгт"),
 pro_y=("Pro жилийн үнэ, ₮", 79000, "plans"),
 pro_mix=("Pro — жилээр төлөгчийн хувь", 0.5, "зорилт ✱"),
 seat_m=("Team суудал сарын үнэ, ₮", 5000, "plans"),
 seat_y=("Team суудал жилийн үнэ, ₮", 50000, "plans"),
 seat_mix=("Team — жилээр төлөгчийн хувь", 0.7, "зорилт ✱"),
 vat=("НӨАТ (үнэд багтсан)", 0.10, "НӨАТ-ын хууль; 50 сая₮-өөс дээш заавал, сайн дураар 10 сая₮-өөс"),
 qpay=("QPay шимтгэл (дотоод карт/банк)", 0.01, "TDB merchant нөхцөл ✱ — гэрээгээр батал"),
 cit=("ААНОАТ (орлогын ≤300 сая₮ бол 1%)", 0.01, "Конcерватив: цэвэр орлогоос тооцов ✱"),
 var_pro=("Хувьсах зардал / Pro хэрэглэгч / сар (AI, имэйл, storage), ₮", 500, "✱"),
 var_seat=("Хувьсах зардал / Team суудал / сар, ₮", 300, "✱"),
 store=("Домэйн + Apple/Google (сард хуваасан), ₮", 50000, "COSTS.md"),
 acct=("Нягтлан (гэрээт), ₮/сар", 300000, "✱"),
 mkt=("Маркетинг, ₮/сар", 2050000, "MARKETING_PLAN §5"),
 payroll=("Цалин (ажилтан, НДШ орсон), ₮/сар", 0, "Үүсгэн байгуулагч цалингүй гэж тооцов. Ажилтан авбал: цалин × 1.125–1.145"),
 scen=("Сценарийн коэффициент (0.5 / 1.0 / 1.5)", 1.0, "Хэрэглэгчийн тоог үржүүлнэ"),
)
keys=list(A)
months=list(range(1,13))
def interp(points):
    out=[]
    for m in months:
        for (m0,v0),(m1,v1) in zip(points,points[1:]):
            if m0<=m<=m1:
                out.append(round(v0+(v1-v0)*(m-m0)/(m1-m0))); break
    return out
pro=interp([(0,0),(3,100),(6,350),(12,900)])
seats=interp([(0,0),(3,60),(6,250),(12,800)])
infra=[200000]*3+[400000]*5+[600000]*4

def compute(k):
    a={x:A[x][1] for x in keys}
    pro_arpa=(1-a['pro_mix'])*a['pro_m']+a['pro_mix']*a['pro_y']/12
    seat_arpa=(1-a['seat_mix'])*a['seat_m']+a['seat_mix']*a['seat_y']/12
    rows=[];cum=0;minc=0
    for i,m in enumerate(months):
        p=pro[i]*k; s=seats[i]*k
        gross=p*pro_arpa+s*seat_arpa
        net=gross/(1+a['vat']); vat=gross-net
        q=gross*a['qpay']; var=p*a['var_pro']+s*a['var_seat']
        gp=net-q-var-infra[i]
        opex=a['store']+a['acct']+a['mkt']+a['payroll']
        ebit=gp-opex; tax=net*a['cit']; ni=ebit-tax
        cum+=ni; minc=min(minc,cum)
        rows.append((m,p,s,gross,vat,net,q,var,infra[i],gp,opex,ebit,tax,ni,cum))
    return rows,minc,pro_arpa,seat_arpa

if __name__=='__main__' and len(sys.argv)>1 and sys.argv[1]=='print':
    for k in (0.5,1.0,1.5):
        rows,minc,pa,sa=compute(k)
        print(f"\n== k={k} pro_arpa={pa:.0f} seat_arpa={sa:.0f} maxcash={minc/1e6:.2f}M yearNI={sum(r[13] for r in rows)/1e6:.2f}M yearGross={sum(r[3] for r in rows)/1e6:.2f}M yearNet={sum(r[5] for r in rows)/1e6:.2f}M")
        be=next((r[0] for r in rows if r[13]>0),None); print('first profitable month',be)
        for r in rows:
            print(r[0],int(r[1]),int(r[2]),*[f"{x/1e3:,.0f}" for x in r[3:]])
    sys.exit()

# ---------- XLSX with live formulas ----------
wb=Workbook(); ws=wb.active; ws.title="Таамаглал"
H=Font(bold=True,color="FFFFFF"); HF=PatternFill("solid",fgColor="1F4E79"); IN=PatternFill("solid",fgColor="FFF2CC")
ws.append(["Таамаглал (шар нүдийг өөрчил)","Утга","Тайлбар / эх сурвалж"])
for c in ws[1]: c.font=H; c.fill=HF
ref={}
for i,k in enumerate(keys,start=2):
    lab,v,n=A[k]; ws.append([lab,v,n]); ws.cell(i,2).fill=IN; ref[k]=f"Таамаглал!$B${i}"
    if isinstance(v,float) and v<1: ws.cell(i,2).number_format='0.0%'
    else: ws.cell(i,2).number_format='#,##0.00' if k=='scen' else '#,##0'
r=len(keys)+3
ws.cell(r,1,"Pro ARPA (НӨАТ орсон), ₮").font=Font(bold=True)
ws.cell(r,2,f"=(1-{ref['pro_mix']})*{ref['pro_m']}+{ref['pro_mix']}*{ref['pro_y']}/12").number_format='#,##0'
ref['pro_arpa']=f"Таамаглал!$B${r}"
ws.cell(r+1,1,"Team суудлын ARPA, ₮").font=Font(bold=True)
ws.cell(r+1,2,f"=(1-{ref['seat_mix']})*{ref['seat_m']}+{ref['seat_mix']}*{ref['seat_y']}/12").number_format='#,##0'
ref['seat_arpa']=f"Таамаглал!$B${r+1}"
ws.cell(r+3,1,"Ханш: 1 USD = 3,595.56₮ (Монголбанк, 2026-09-25). ✱ = таамаг — бодит өгөгдлөөр сар бүр шинэчил.")
ws.column_dimensions['A'].width=58; ws.column_dimensions['B'].width=16; ws.column_dimensions['C'].width=70

m=wb.create_sheet("12 сар (P&L)")
labels=[("Төлбөртэй Pro хэрэглэгч (төлөвлөгөө)",'in_pro'),("Team суудал (төлөвлөгөө)",'in_seat'),
("Pro × сценари",'p'),("Суудал × сценари",'s'),
("Нийт борлуулалт (НӨАТ орсон)",'gross'),("НӨАТ (төсөвт төлнө)",'vat'),("Цэвэр орлого",'net'),
("QPay шимтгэл",'q'),("Хувьсах зардал (AI, имэйл)",'var'),("Дэд бүтэц (Supabase, Netlify, Resend…)",'infra'),
("Нийт ашиг (gross profit)",'gp'),("Gross margin %",'gm'),
("Домэйн + store",'store'),("Нягтлан",'acct'),("Маркетинг",'mkt'),("Цалин",'payroll'),("Үйл ажиллагааны зардал",'opex'),
("Татварын өмнөх ашиг",'ebit'),("ААНОАТ",'tax'),("ЦЭВЭР АШИГ",'ni'),("Хуримтлагдсан ашиг (мөнгөний хэрэгцээ)",'cum')]
m.append(["Мөр / Сар"]+[f"{i}-р сар" for i in months]+["12 сар нийт"])
for c in m[1]: c.font=H; c.fill=HF
row={k:i+2 for i,(_,k) in enumerate(labels)}
for i,(lab,k) in enumerate(labels):
    rr=i+2; m.cell(rr,1,lab)
    for j,mm in enumerate(months):
        col=L(j+2); R=lambda key: f"{col}{row[key]}"
        f={'in_pro':pro[j],'in_seat':seats[j],
           'p':f"={R('in_pro')}*{ref['scen']}",'s':f"={R('in_seat')}*{ref['scen']}",
           'gross':f"={R('p')}*{ref['pro_arpa']}+{R('s')}*{ref['seat_arpa']}",
           'vat':f"={R('gross')}-{R('net')}",'net':f"={R('gross')}/(1+{ref['vat']})",
           'q':f"={R('gross')}*{ref['qpay']}",'var':f"={R('p')}*{ref['var_pro']}+{R('s')}*{ref['var_seat']}",
           'infra':infra[j],'gp':f"={R('net')}-{R('q')}-{R('var')}-{R('infra')}",
           'gm':f"=IF({R('net')}=0,0,{R('gp')}/{R('net')})",
           'store':f"={ref['store']}",'acct':f"={ref['acct']}",'mkt':f"={ref['mkt']}",'payroll':f"={ref['payroll']}",
           'opex':f"={R('store')}+{R('acct')}+{R('mkt')}+{R('payroll')}",
           'ebit':f"={R('gp')}-{R('opex')}",'tax':f"={R('net')}*{ref['cit']}",'ni':f"={R('ebit')}-{R('tax')}",
           'cum':f"={R('ni')}" if j==0 else f"={L(j+1)}{row['cum']}+{R('ni')}"}[k]
        c=m.cell(rr,j+2,f); c.number_format='0%' if k=='gm' else '#,##0'
        if k in('in_pro','in_seat','infra'): c.fill=IN
    tot=m.cell(rr,14)
    if k in('gross','vat','net','q','var','infra','gp','store','acct','mkt','payroll','opex','ebit','tax','ni'):
        tot.value=f"=SUM(B{rr}:M{rr})"; tot.number_format='#,##0'
    elif k=='cum': tot.value=f"=MIN(B{rr}:M{rr})"; tot.number_format='#,##0'; m.cell(rr,15,"← хамгийн бага = шаардлагатай мөнгө")
    if k in('gross','gp','ni','cum'):
        for c in m[rr]: c.font=Font(bold=True)
m.column_dimensions['A'].width=44
for j in range(2,15): m.column_dimensions[L(j)].width=13
m.freeze_panes="B2"

s=wb.create_sheet("Сценари")
s.append(["Сценари","Коэф.","12 сарын борлуулалт (НӨАТ орсон)","12 сарын цэвэр ашиг","Мөнгөний хэрэгцээ (хамгийн бага хуримтлал)","Анхны ашигтай сар"])
for c in s[1]: c.font=H; c.fill=HF
for name,k in (("Болгоомжтой",0.5),("Суурь",1.0),("Өөдрөг",1.5)):
    rows,minc,_,_=compute(k)
    be=next((r[0] for r in rows if r[13]>0),"12 сард хүрэхгүй")
    s.append([name,k,round(sum(r[3] for r in rows)),round(sum(r[13] for r in rows)),round(minc),be])
for rr in range(2,5):
    for cc in (3,4,5): s.cell(rr,cc).number_format='#,##0'
s.append([]); s.append(["Энэ хуудас нь скриптээр тооцсон тогтмол утга. Өөр утга авах бол «Таамаглал»-ын коэффициентыг өөрчилж «12 сар (P&L)»-г хар."])
for col,w in zip("ABCDEF",(14,8,32,22,40,20)): s.column_dimensions[col].width=w
wb.save(sys.argv[1] if len(sys.argv)>1 else "financial_model.xlsx")
print("saved")
